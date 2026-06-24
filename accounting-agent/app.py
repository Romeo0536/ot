"""Web interface — อัปโหลดบิล → AI อ่าน → ดาวน์โหลด Excel.

วิธีรัน:
    streamlit run app.py
หรือดับเบิลคลิกที่ run_app.bat (Windows)
"""
from __future__ import annotations

import io
import os
import tempfile
import time
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv

# โหลด .env ก่อนสิ่งอื่น (ต้องมี GEMINI_API_KEY อยู่ใน .env)
load_dotenv()

# --------------------------------------------------------------------------- #
# Page config
# --------------------------------------------------------------------------- #
st.set_page_config(
    page_title="AI Agent นักบัญชี",
    page_icon="🧾",
    layout="wide",
)

# --------------------------------------------------------------------------- #
# Load config (อ่านก่อนเพื่อรู้ว่าใช้ provider ไหน)
# --------------------------------------------------------------------------- #
@st.cache_resource(show_spinner=False)
def _load_config():
    from accounting_agent.config import Config

    config_path = Path(__file__).parent / "config.yaml"
    return Config.load(config_path)


_cfg = _load_config()

# ตัวเลือก provider / model ใน UI (key = ชื่อใน config, label = แสดงในเมนู)
_PROVIDERS = {
    "gemini": "🔮 Gemini (Google) — ฟรี เร็ว แม่น",
    "ollama": "🦙 Ollama (local) — ฟรี 100% ออฟไลน์",
    "groq": "⚡ Groq — ฟรี เร็วมาก (ไฟล์เล็ก)",
    "claude": "🤖 Claude (Anthropic) — เสถียร เสียเงิน",
}

# โมเดลแนะนำของแต่ละ provider — เลือกจาก dropdown หรือพิมพ์เองได้
_MODELS = {
    "gemini": ["gemini-2.5-flash", "gemini-2.5-pro", "gemini-2.0-flash"],
    "ollama": ["qwen2.5vl:7b", "llava:7b", "minicpm-v", "llama3.2", "mistral"],
    "groq": [
        "meta-llama/llama-4-scout-17b-16e-instruct",  # vision — อ่านรูปได้
        "meta-llama/llama-4-maverick-17b-128e-instruct",
        "llama-3.3-70b-versatile",  # text only
    ],
    "claude": ["claude-haiku-4-5", "claude-opus-4-8", "claude-sonnet-4-6"],
}

# provider ที่ต้องใช้ API key (Ollama รัน local ไม่ต้องมี key)
_API_KEY_ENV = {
    "gemini": "GEMINI_API_KEY",
    "claude": "ANTHROPIC_API_KEY",
    "groq": "GROQ_API_KEY",
}
_PROVIDER_LABEL = {
    "gemini": "Gemini API Key",
    "claude": "Anthropic API Key",
    "groq": "Groq API Key",
}
_KEY_HELP = {
    "gemini": "ขอ key ฟรีที่ https://aistudio.google.com/apikey",
    "claude": "ขอ key ที่ https://console.anthropic.com/",
    "groq": "ขอ key ฟรีที่ https://console.groq.com/",
}

# ค่าตั้งต้นจาก config.yaml — ใช้ครั้งแรกเท่านั้น (หลังจากนั้น session_state เก็บค่าผู้ใช้เลือก)
if "provider" not in st.session_state:
    st.session_state["provider"] = (
        _cfg.provider.lower() if _cfg.provider.lower() in _PROVIDERS else "gemini"
    )
if "model" not in st.session_state:
    st.session_state["model"] = _cfg.model

# --------------------------------------------------------------------------- #
# Sidebar — เลือก AI provider + model + API key
# --------------------------------------------------------------------------- #
with st.sidebar:
    st.header("⚙️ ตั้งค่า")

    # 1) เลือก provider
    provider_keys = list(_PROVIDERS.keys())
    selected_provider = st.selectbox(
        "AI Provider",
        options=provider_keys,
        index=provider_keys.index(st.session_state["provider"]),
        format_func=lambda k: _PROVIDERS[k],
    )
    # ถ้าผู้ใช้เปลี่ยน provider → รีเซ็ต model เป็นตัวแรกของ provider ใหม่
    if selected_provider != st.session_state["provider"]:
        st.session_state["provider"] = selected_provider
        st.session_state["model"] = _MODELS[selected_provider][0]
        st.rerun()

    _PROVIDER = st.session_state["provider"]

    # 2) เลือก model (พิมพ์เองได้)
    model_options = _MODELS[_PROVIDER]
    current_model = st.session_state["model"]
    if current_model not in model_options:
        model_options = [current_model] + model_options
    selected_model = st.selectbox(
        "Model",
        options=model_options,
        index=model_options.index(current_model),
        accept_new_options=True,
        help="พิมพ์ชื่อ model อื่นได้ ถ้าไม่มีในรายการ",
    )
    st.session_state["model"] = selected_model

    st.divider()

    # 3) API key (ถ้า provider ต้องใช้)
    needs_key = _PROVIDER in _API_KEY_ENV
    api_key_input = ""
    if needs_key:
        env_name = _API_KEY_ENV[_PROVIDER]
        api_key_input = st.text_input(
            _PROVIDER_LABEL[_PROVIDER],
            value=os.getenv(env_name, ""),
            type="password",
            help=_KEY_HELP[_PROVIDER],
            placeholder="วาง API key ที่นี่",
        )
        if api_key_input:
            os.environ[env_name] = api_key_input
    else:
        st.success("🦙 ใช้ Ollama (local) — ไม่ต้องใช้ API key", icon="✅")
        st.caption("ต้องเปิดโปรแกรม Ollama ค้างไว้ (ไอคอนลามะมุมขวาล่าง)")

    st.divider()
    if needs_key:
        st.caption(
            "**วิธีใช้**\n"
            "1. วาง API key ด้านบน\n"
            "2. อัปโหลด PDF / รูปใบเสร็จ\n"
            "3. กด **เริ่มอ่าน**\n"
            "4. ดาวน์โหลด Excel"
        )
    else:
        st.caption(
            "**วิธีใช้**\n"
            "1. เปิดโปรแกรม Ollama\n"
            "2. อัปโหลด PDF / รูปใบเสร็จ\n"
            "3. กด **เริ่มอ่าน**\n"
            "4. ดาวน์โหลด Excel"
        )
    st.caption("รองรับ: PDF, JPG, PNG, WEBP, GIF")

# --------------------------------------------------------------------------- #
# Header
# --------------------------------------------------------------------------- #
st.title("🧾 AI Agent นักบัญชี")
st.caption("อัปโหลดบิล/ใบเสร็จ (PDF หรือรูปภาพ) → AI อ่านข้อมูล → ดาวน์โหลด Excel")


def _get_provider():
    from accounting_agent.providers import get_provider

    return get_provider(st.session_state["provider"], st.session_state["model"])


# --------------------------------------------------------------------------- #
# File uploader
# --------------------------------------------------------------------------- #
uploaded_files = st.file_uploader(
    "เลือกหรือลากไฟล์ใบเสร็จมาวางที่นี่",
    type=["pdf", "jpg", "jpeg", "png", "webp", "gif"],
    accept_multiple_files=True,
    label_visibility="collapsed",
)

if not uploaded_files:
    st.info("⬆️  อัปโหลดไฟล์ด้านบนเพื่อเริ่มต้น", icon="📄")

# --------------------------------------------------------------------------- #
# Process button
# --------------------------------------------------------------------------- #
if uploaded_files:
    n = len(uploaded_files)
    btn_label = f"🚀 เริ่มอ่านใบเสร็จ ({n} ไฟล์)"
    run = st.button(btn_label, type="primary", use_container_width=True)

    if run:
        # provider ที่ต้องใช้ key — เช็คก่อนว่าใส่แล้วยัง
        if needs_key and not os.getenv(_API_KEY_ENV[_PROVIDER], ""):
            st.error(f"กรุณาใส่ {_PROVIDER_LABEL[_PROVIDER]} ในแถบซ้าย", icon="🔑")
            st.stop()

        config = _cfg
        try:
            provider = _get_provider()
        except Exception as e:
            st.error(f"เชื่อมต่อ AI ไม่ได้: {e}", icon="❌")
            st.stop()

        from accounting_agent.extract import extract_receipts
        from accounting_agent.store import thai_date

        all_rows: list[dict] = []
        errors: list[str] = []

        progress_bar = st.progress(0, text="กำลังเริ่มต้น...")
        log_area = st.empty()
        log_lines: list[str] = []

        with tempfile.TemporaryDirectory() as tmpdir:
            for i, uf in enumerate(uploaded_files):
                log_lines.append(f"⏳ กำลังอ่าน **{uf.name}**…")
                log_area.markdown("\n\n".join(log_lines))
                progress_bar.progress((i) / n, text=f"อ่านไฟล์ {i+1}/{n}: {uf.name}")

                tmp_path = Path(tmpdir) / uf.name
                tmp_path.write_bytes(uf.getvalue())

                try:
                    receipts = extract_receipts(tmp_path, config, provider)
                    for r in receipts:
                        all_rows.append(
                            {
                                "วันที่": r.receipt_date or "",
                                "วันที่ (พ.ศ.)": thai_date(r.receipt_date),
                                "ร้านค้า": r.vendor,
                                "หมวด": r.category,
                                "ยอดรวม": r.total_amount,
                                "VAT": r.tax_amount if r.tax_amount is not None else "",
                                "สกุลเงิน": r.currency,
                                "เลขที่ผู้เสียภาษี": r.tax_id or "",
                                "เลขที่ใบเสร็จ": r.invoice_number or "",
                                "วิธีชำระ": r.payment_method or "",
                                "ความมั่นใจ": r.confidence,
                                "ไฟล์ต้นฉบับ": uf.name,
                                "หมายเหตุ": r.notes or "",
                            }
                        )
                    flag = "  ⚠️" if any(r.confidence < 0.6 for r in receipts) else ""
                    log_lines[-1] = (
                        f"✅ **{uf.name}** — พบ {len(receipts)} ใบเสร็จ{flag}"
                    )
                except Exception as exc:
                    log_lines[-1] = f"❌ **{uf.name}** — {exc}"
                    errors.append(f"{uf.name}: {exc}")

                log_area.markdown("\n\n".join(log_lines))
                progress_bar.progress((i + 1) / n, text=f"เสร็จ {i+1}/{n}")

        progress_bar.empty()

        if all_rows:
            st.session_state["results"] = all_rows
        if not all_rows and errors:
            hint = "ตรวจสอบว่าเปิด Ollama แล้ว" if not needs_key else "ตรวจสอบ API Key"
            st.error(f"อ่านไม่ได้เลย — {hint} และลองใหม่อีกครั้ง", icon="❌")

# --------------------------------------------------------------------------- #
# Results table & Excel download
# --------------------------------------------------------------------------- #
if "results" in st.session_state and st.session_state["results"]:
    import pandas as pd

    rows = st.session_state["results"]
    df = pd.DataFrame(rows)

    st.divider()
    st.subheader(f"📋 ผลการอ่าน ({len(df)} รายการ)")

    # ไฮไลต์แถวที่ confidence ต่ำ
    def _highlight_low(row):
        if row.get("ความมั่นใจ", 1.0) < 0.6:
            return ["background-color: #fff3cd"] * len(row)
        return [""] * len(row)

    st.dataframe(
        df.style.apply(_highlight_low, axis=1),
        use_container_width=True,
        hide_index=True,
        column_config={
            "ยอดรวม": st.column_config.NumberColumn(format="%.2f"),
            "VAT": st.column_config.NumberColumn(format="%.2f"),
            "ความมั่นใจ": st.column_config.ProgressColumn(
                min_value=0, max_value=1, format="%.0%%"
            ),
        },
    )

    if any(r.get("ความมั่นใจ", 1.0) < 0.6 for r in rows):
        st.warning("⚠️ แถวสีเหลืองมีความมั่นใจต่ำ — แนะนำให้ตรวจสอบซ้ำ")

    # สร้าง Excel
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="ใบเสร็จ", index=False)
        ws = writer.sheets["ใบเสร็จ"]

        # ตั้งความกว้างคอลัมน์อัตโนมัติ
        for col_cells in ws.columns:
            col_letter = col_cells[0].column_letter
            max_len = max(
                (len(str(cell.value)) if cell.value is not None else 0)
                for cell in col_cells
            )
            ws.column_dimensions[col_letter].width = min(max_len + 3, 45)

        # ทำ header เป็นตัวหนา
        from openpyxl.styles import Font, PatternFill, Alignment

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # freeze header row
        ws.freeze_panes = "A2"

    st.download_button(
        label="⬇️ ดาวน์โหลด Excel (.xlsx)",
        data=buf.getvalue(),
        file_name="ledger.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        type="primary",
    )

    if st.button("🗑️ ล้างผลลัพธ์", use_container_width=True):
        del st.session_state["results"]
        st.rerun()
